"""
외국인 학과시험 AI 튜터 — 혁신성과경진대회 PPT 자동 생성

사용법:
    python scripts/build_report_ppt.py [--out PATH]

요구사항:
    pip install python-pptx (이미 설치됨)

출력:
    기본: C:/Users/user/OneDrive/바탕 화면/외국인_학과시험_AI튜터_보고서_YYYY-MM.pptx
    --out 옵션으로 다른 경로 지정 가능

설계:
    바탕화면의 '2025년 5월 2주차 혁신성장 전략회의' PDF 양식 재현
    16:9 와이드 / 파란색 #187DFA 강조 / 헤더 형식 "01 추진배경 • 세부"
    15장 슬라이드 자동 빌드
"""

import argparse
import os
import sys
import io
from datetime import datetime
from pathlib import Path

# 콘솔 출력 UTF-8 강제 (PowerShell cp949 인코딩 회피)
try:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', line_buffering=True)
except Exception:
    pass

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree


# ----- 색상·폰트 상수 -----

C_PRIMARY = RGBColor(0x18, 0x7D, 0xFA)        # 메인 파랑 (헤더·강조)
C_DEEP = RGBColor(0x00, 0x3F, 0xA2)            # 짙은 파랑 (제목)
C_ACCENT = RGBColor(0x98, 0x5E, 0xE3)          # 보라 (수치 강조)
C_SUCCESS = RGBColor(0x00, 0x92, 0x3F)         # 초록 (성과)
C_TEXT = RGBColor(0x26, 0x26, 0x26)            # 본문 텍스트
C_GRAY = RGBColor(0x7F, 0x7F, 0x7F)            # 회색 (보조)
C_LIGHT_BG = RGBColor(0xF3, 0xF4, 0xF8)        # 연회색 배경
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK = RGBColor(0x00, 0x00, 0x00)

FONT_KO = "맑은 고딕"
FONT_EN = "Calibri"


# ----- 자료 경로 -----

PROJECT_ROOT = Path(__file__).parent.parent.resolve()
LOGO_PATH = Path(r"C:\Users\user\OneDrive\바탕 화면\CI편집\2. 부산지부\시도지부 결합 A -1.png")
QR_PATH = PROJECT_ROOT / "외국인면허취득 QR.png"


# ----- 헬퍼 -----

def add_textbox(slide, left, top, width, height, text, *,
                font_size=14, bold=False, color=C_TEXT, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP, font_name=FONT_KO):
    """텍스트박스 추가"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def add_rect(slide, left, top, width, height, fill_color=None, line_color=None,
             line_width=0.5, no_fill=False):
    """사각형 도형 추가"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if no_fill:
        shape.fill.background()
    elif fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    if line_color is not None:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_rounded(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    if line_color is not None:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def set_shape_text(shape, text, *, font_size=14, bold=False, color=C_WHITE,
                   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font_name=FONT_KO):
    tf = shape.text_frame
    tf.margin_left = Emu(50000)
    tf.margin_right = Emu(50000)
    tf.margin_top = Emu(20000)
    tf.margin_bottom = Emu(20000)
    tf.vertical_anchor = anchor
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_header(slide, section_label, sub_label, page_num, total_pages):
    """공통 페이지 헤더: 좌상단 섹션 + 우하단 페이지 번호"""
    # 좌상단 섹션 라벨
    add_textbox(slide, Inches(0.4), Inches(0.25), Inches(11), Inches(0.4),
                f"{section_label}  ●  {sub_label}",
                font_size=14, bold=True, color=C_PRIMARY)
    # 헤더 하단 줄
    line = slide.shapes.add_connector(1, Inches(0.4), Inches(0.7), Inches(13), Inches(0.7))
    line.line.color.rgb = C_PRIMARY
    line.line.width = Pt(2)
    # 우하단 페이지 번호
    add_textbox(slide, Inches(12.5), Inches(7.0), Inches(0.7), Inches(0.4),
                f"{page_num} / {total_pages}",
                font_size=10, color=C_GRAY, align=PP_ALIGN.RIGHT)


def add_logo(slide, left, top, width):
    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), left, top, width=width)


# ----- 슬라이드 빌더 -----

TOTAL_PAGES = 15


def build_cover(prs):
    """슬라이드 1 — 표지"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # 배경 강조 도형 (좌측 세로 띠)
    add_rect(slide, Inches(0), Inches(0), Inches(0.6), Inches(7.5), fill_color=C_DEEP)
    add_rect(slide, Inches(0.6), Inches(0), Inches(0.15), Inches(7.5), fill_color=C_PRIMARY)

    # 로고 (오른쪽 상단)
    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(10.5), Inches(0.4), width=Inches(2.5))

    # 부서명
    add_textbox(slide, Inches(1.5), Inches(1.0), Inches(11), Inches(0.5),
                "한국도로교통공단 부산광역시지부 안전교육부",
                font_size=16, bold=True, color=C_GRAY)

    # 메인 제목
    add_textbox(slide, Inches(1.5), Inches(2.3), Inches(11), Inches(1.5),
                "외국인 학과시험 AI 튜터",
                font_size=54, bold=True, color=C_DEEP)

    # 부제
    add_textbox(slide, Inches(1.5), Inches(3.6), Inches(11), Inches(0.8),
                "AI · 디지털 활용으로 모국어 면허 학습 환경 구축",
                font_size=22, color=C_PRIMARY)

    # 라벨 박스 (강조)
    box = add_rounded(slide, Inches(1.5), Inches(4.8), Inches(6.5), Inches(1.0),
                       fill_color=C_PRIMARY)
    set_shape_text(box,
                   "GPTs AI 튜터  ·  외국인 맞춤 교재  ·  4언어 1000문항  ·  자동 갱신",
                   font_size=15, bold=True, color=C_WHITE)

    # 일자
    add_textbox(slide, Inches(1.5), Inches(6.3), Inches(11), Inches(0.5),
                f"2026. 5.",
                font_size=18, color=C_TEXT)

    return slide


def build_overview(prs):
    """슬라이드 2 — 기본현황 / 4대 통합 솔루션"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "기본현황", "외국인 학과시험 AI 튜터 4대 통합 솔루션", 2, TOTAL_PAGES)

    # 메인 메시지
    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.7),
                "한국 거주 외국인의 모국어 면허 학습을 위한",
                font_size=20, color=C_TEXT, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.5), Inches(1.5), Inches(12.5), Inches(0.7),
                "AI · 디지털 통합 솔루션",
                font_size=28, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    # 4개 카드 (가로 배치)
    cards = [
        ("🤖", "GPTs AI 튜터", "ChatGPT 기반 챗봇", "모국어 질문 → AI 응답·해설", C_PRIMARY),
        ("📘", "외국인 맞춤 교재", "PC · 모바일 반응형", "4개 언어 체계적 학습", C_DEEP),
        ("🌐", "1000문항 학습 사이트", "한·영·중·베", "동영상 35개 포함", C_ACCENT),
        ("🔄", "자동 갱신 시스템", "PDF → HTML 자동", "지속 운영 가능", C_SUCCESS),
    ]
    card_w = Inches(2.95)
    card_gap = Inches(0.2)
    start_x = Inches(0.5)
    card_y = Inches(2.7)
    card_h = Inches(3.3)

    for i, (icon, title, sub1, sub2, color) in enumerate(cards):
        x = start_x + (card_w + card_gap) * i
        # 카드 배경
        add_rounded(slide, x, card_y, card_w, card_h, fill_color=C_LIGHT_BG)
        # 컬러 헤더 띠
        add_rect(slide, x, card_y, card_w, Inches(0.5), fill_color=color)
        # 아이콘
        add_textbox(slide, x, card_y + Inches(0.7), card_w, Inches(0.8),
                    icon, font_size=40, align=PP_ALIGN.CENTER)
        # 제목
        add_textbox(slide, x, card_y + Inches(1.7), card_w, Inches(0.5),
                    title, font_size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
        # 부제 1
        add_textbox(slide, x, card_y + Inches(2.3), card_w, Inches(0.4),
                    sub1, font_size=12, color=C_TEXT, align=PP_ALIGN.CENTER)
        # 부제 2
        add_textbox(slide, x, card_y + Inches(2.7), card_w, Inches(0.4),
                    sub2, font_size=12, color=C_GRAY, align=PP_ALIGN.CENTER)

    # 추진 단계 안내 (하단)
    add_textbox(slide, Inches(0.5), Inches(6.3), Inches(12.5), Inches(0.5),
                "추진 단계 :  추진배경  →  추진계획  →  추진내용 및 결과  →  향후계획",
                font_size=14, color=C_GRAY, align=PP_ALIGN.CENTER)


def build_bg_1(prs):
    """슬라이드 3 — 01 추진배경 ① 한국 거주 외국인 증가"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "01 추진배경", "한국 거주 외국인 증가 + 면허 취득 수요", 3, TOTAL_PAGES)

    # 큰 헤드라인
    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.7),
                "한국의 위상이 높아짐에 따라",
                font_size=22, color=C_TEXT, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.5), Inches(1.6), Inches(12.5), Inches(0.9),
                "국내 외국인 거주자가 빠르게 늘어나고 있다",
                font_size=30, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    # 핵심 통계 박스 3개 (가로)
    stats = [
        ("250만 명+", "법무부 통계 기준\n국내 체류 외국인 (2024)", C_PRIMARY),
        ("증가 추세", "사회 · 경제 활동 외국인\n지속 확대", C_ACCENT),
        ("면허 수요 ↑", "직장 · 가족 · 일상\n자가운전 필수 인프라", C_SUCCESS),
    ]
    box_w = Inches(3.8)
    box_gap = Inches(0.4)
    start_x = Inches(0.85)
    box_y = Inches(3.0)
    box_h = Inches(2.0)

    for i, (big, small, color) in enumerate(stats):
        x = start_x + (box_w + box_gap) * i
        add_rounded(slide, x, box_y, box_w, box_h, fill_color=C_LIGHT_BG)
        add_textbox(slide, x, box_y + Inches(0.3), box_w, Inches(0.8),
                    big, font_size=28, bold=True, color=color, align=PP_ALIGN.CENTER)
        add_textbox(slide, x, box_y + Inches(1.15), box_w, Inches(0.8),
                    small, font_size=12, color=C_TEXT, align=PP_ALIGN.CENTER)

    # 마무리 메시지
    add_textbox(slide, Inches(0.5), Inches(5.7), Inches(12.5), Inches(0.7),
                "→ 외국인을 위한 면허 취득 학습 환경이 시급한 사회적 과제",
                font_size=18, bold=True, color=C_PRIMARY, align=PP_ALIGN.CENTER)


def build_bg_2(prs):
    """슬라이드 4 — 01 추진배경 ② 외국인 면허 학습의 3가지 어려움"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "01 추진배경", "외국인 면허 학습의 3가지 어려움", 4, TOTAL_PAGES)

    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.8),
                "기존 환경에서 외국인이 마주치는 벽",
                font_size=26, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    # 3가지 문제 카드 (세로 배치)
    problems = [
        ("①", "모국어 정보 부족",
         "도로교통법 · 면허 취득 절차 등 기본 정보가\n외국인의 모국어로 제공되지 않음"),
        ("②", "외국인 맞춤 교재 부재",
         "단순 번역본은 있어도, 외국인 학습자 관점에서\n재구성된 체계적 교재가 없음"),
        ("③", "학과시험 4개 언어 + 해설 없음",
         "한 · 영 · 중 · 베 4개 언어만 지원,\n그마저도 해설 없이 정답만 제공"),
    ]
    item_y = Inches(2.2)
    item_h = Inches(1.3)
    item_gap = Inches(0.2)

    for i, (no, title, desc) in enumerate(problems):
        y = item_y + (item_h + item_gap) * i
        # 번호 박스 (좌측)
        no_box = add_rounded(slide, Inches(0.7), y, Inches(1.3), item_h,
                              fill_color=C_PRIMARY)
        set_shape_text(no_box, no, font_size=44, bold=True, color=C_WHITE)
        # 본문 박스 (우측)
        add_rounded(slide, Inches(2.2), y, Inches(10.5), item_h, fill_color=C_LIGHT_BG)
        add_textbox(slide, Inches(2.5), y + Inches(0.18), Inches(10), Inches(0.5),
                    title, font_size=20, bold=True, color=C_DEEP)
        add_textbox(slide, Inches(2.5), y + Inches(0.7), Inches(10), Inches(0.6),
                    desc, font_size=13, color=C_TEXT)


def build_bg_3(prs):
    """슬라이드 5 — 01 추진배경 ③ 공공 서비스 형평성"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "01 추진배경", "공공 서비스 형평성 · 외국인 정착 지원", 5, TOTAL_PAGES)

    # 큰 인용 박스
    quote_box = add_rounded(slide, Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.3),
                             fill_color=C_DEEP)
    set_shape_text(quote_box,
                   '"면허는 외국인 정착의 기본 인프라"',
                   font_size=28, bold=True, color=C_WHITE)

    # 핵심 메시지 3가지
    msgs = [
        ("🌏", "사회적 가치", "한국 거주 외국인의 일상생활 · 직장 활동을 안전하게"),
        ("⚖️", "공공 서비스 형평성", "내 · 외국인 모두에게 동등한 학습 접근권 보장"),
        ("🏛️", "공단의 사회적 책임", "교통안전 향상 + 외국인 정착 지원이라는 공익 목적"),
    ]
    msg_y = Inches(3.1)
    msg_h = Inches(1.0)
    msg_gap = Inches(0.15)

    for i, (icon, title, desc) in enumerate(msgs):
        y = msg_y + (msg_h + msg_gap) * i
        # 아이콘 박스
        add_rounded(slide, Inches(0.8), y, Inches(1.0), msg_h, fill_color=C_PRIMARY)
        add_textbox(slide, Inches(0.8), y + Inches(0.15), Inches(1.0), Inches(0.7),
                    icon, font_size=32, align=PP_ALIGN.CENTER)
        # 본문
        add_textbox(slide, Inches(2.0), y + Inches(0.1), Inches(10.5), Inches(0.45),
                    title, font_size=18, bold=True, color=C_DEEP)
        add_textbox(slide, Inches(2.0), y + Inches(0.55), Inches(10.5), Inches(0.45),
                    desc, font_size=14, color=C_TEXT)

    # 결론
    add_textbox(slide, Inches(0.5), Inches(6.55), Inches(12.5), Inches(0.5),
                "→ 외국인도 모국어로 안전하게 면허를 학습할 수 있어야 한다",
                font_size=18, bold=True, color=C_PRIMARY, align=PP_ALIGN.CENTER)


def build_plan(prs):
    """슬라이드 6 — 02 추진계획 AI · 디지털 활용 4대 축"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "02 추진계획", "AI · 디지털 활용 4대 축 통합 솔루션", 6, TOTAL_PAGES)

    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.6),
                "단순 자료 배포가 아닌 AI 기반 통합 학습 환경 구축",
                font_size=20, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    # 중심 학습자
    center_box = add_rounded(slide, Inches(5.5), Inches(3.5), Inches(2.3), Inches(1.0),
                              fill_color=C_DEEP)
    set_shape_text(center_box, "외국인 학습자\n(모국어 사용)",
                    font_size=16, bold=True, color=C_WHITE)

    # 4개 솔루션 박스 (사방으로 배치)
    solutions = [
        # (left, top, label, sub, color)
        (Inches(0.8), Inches(2.0), "🤖 GPTs AI 튜터", "맞춤 질의응답", C_PRIMARY),
        (Inches(10.4), Inches(2.0), "📘 외국인 맞춤 교재", "체계적 학습", C_ACCENT),
        (Inches(0.8), Inches(5.0), "🌐 학습 사이트", "1000문항 연습", C_SUCCESS),
        (Inches(10.4), Inches(5.0), "🔄 자동 갱신", "지속가능성", RGBColor(0xE6, 0x7E, 0x22)),
    ]
    for left, top, label, sub, color in solutions:
        sol_box = add_rounded(slide, left, top, Inches(2.3), Inches(1.3), fill_color=color)
        tf = sol_box.text_frame
        tf.margin_left = Emu(50000)
        tf.margin_right = Emu(50000)
        tf.margin_top = Emu(50000)
        tf.margin_bottom = Emu(50000)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        run1 = p1.add_run()
        run1.text = label
        run1.font.name = FONT_KO
        run1.font.size = Pt(14)
        run1.font.bold = True
        run1.font.color.rgb = C_WHITE
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = sub
        run2.font.name = FONT_KO
        run2.font.size = Pt(11)
        run2.font.color.rgb = C_WHITE

    # 하단 강조
    add_textbox(slide, Inches(0.5), Inches(6.5), Inches(12.5), Inches(0.5),
                "✓  AI · 다국어 · 자동화 통합으로 어떤 외국인도 자신의 언어로 학습 가능",
                font_size=14, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)


def build_content_1_gpts(prs):
    """슬라이드 7 — 03 추진내용 ① GPTs AI 튜터"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "03 추진내용 ①", "GPTs AI 튜터 (모국어 챗봇)", 7, TOTAL_PAGES)

    # 좌측: 설명
    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(8), Inches(0.7),
                "🤖  ChatGPT 기반 외국인 전용 면허 챗봇",
                font_size=22, bold=True, color=C_PRIMARY)

    items = [
        ("✓", "어떤 언어로든 질문 가능", "외국인이 모국어로 자유롭게 질문"),
        ("✓", "4개 언어 문제은행 학습 데이터", "한 · 영 · 중 · 베 1000문항 + 교재 학습"),
        ("✓", "모국어로 응답 · 해설", "정답 + 자세한 해설을 자기 언어로 제공"),
        ("✓", "24 / 7 운영", "전 세계 어디서든 시간 무관 이용"),
    ]
    for i, (mark, title, desc) in enumerate(items):
        y = Inches(2.0 + i * 0.85)
        # 체크 마크
        add_textbox(slide, Inches(0.5), y, Inches(0.5), Inches(0.5),
                    mark, font_size=20, bold=True, color=C_SUCCESS)
        # 제목
        add_textbox(slide, Inches(1.0), y, Inches(7.5), Inches(0.4),
                    title, font_size=15, bold=True, color=C_DEEP)
        # 설명
        add_textbox(slide, Inches(1.0), y + Inches(0.4), Inches(7.5), Inches(0.4),
                    desc, font_size=12, color=C_GRAY)

    # 우측: QR + URL 박스
    qr_panel = add_rounded(slide, Inches(8.8), Inches(1.5), Inches(4.2), Inches(5.5),
                            fill_color=C_LIGHT_BG)
    add_textbox(slide, Inches(8.8), Inches(1.6), Inches(4.2), Inches(0.5),
                "🔗 GPTs 접속 URL · QR",
                font_size=14, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    if QR_PATH.exists():
        slide.shapes.add_picture(str(QR_PATH), Inches(9.6), Inches(2.2),
                                  width=Inches(2.6), height=Inches(2.6))

    add_textbox(slide, Inches(8.9), Inches(5.0), Inches(4.0), Inches(0.4),
                "Driver's License",
                font_size=12, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(8.9), Inches(5.35), Inches(4.0), Inches(0.4),
                "AI Tutor for Foreigners",
                font_size=12, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(8.9), Inches(5.85), Inches(4.0), Inches(0.4),
                "chatgpt.com/g/...",
                font_size=10, color=C_GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(8.9), Inches(6.25), Inches(4.0), Inches(0.4),
                "QR · URL로 즉시 접속",
                font_size=11, color=C_PRIMARY, align=PP_ALIGN.CENTER)


def build_content_2_textbook(prs):
    """슬라이드 8 — 03 추진내용 ② 외국인 맞춤 교재"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "03 추진내용 ②", "외국인 맞춤 기본 교재 (PC · 모바일)", 8, TOTAL_PAGES)

    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.7),
                "📘  단순 번역이 아닌, 외국인 학습자 관점에서 재구성한 체계적 교재",
                font_size=20, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    # 4개 언어 카드
    langs = [
        ("🇰🇷 한국어", "외국인 운전면허\n학과시험 교재", C_PRIMARY),
        ("🇺🇸 English", "Korean Driver's License\nWritten Test Textbook", C_DEEP),
        ("🇨🇳 中文", "韩国驾驶执照\n笔试教材", C_ACCENT),
        ("🇻🇳 Tiếng Việt", "Giáo trình thi lý thuyết\nbằng lái xe Hàn Quốc", C_SUCCESS),
    ]
    card_w = Inches(2.95)
    card_h = Inches(2.0)
    card_y = Inches(2.2)
    for i, (flag, title, color) in enumerate(langs):
        x = Inches(0.5) + (card_w + Inches(0.2)) * i
        add_rounded(slide, x, card_y, card_w, card_h, fill_color=C_LIGHT_BG)
        add_rect(slide, x, card_y, card_w, Inches(0.5), fill_color=color)
        add_textbox(slide, x, card_y + Inches(0.7), card_w, Inches(0.5),
                    flag, font_size=20, bold=True, color=color, align=PP_ALIGN.CENTER)
        add_textbox(slide, x, card_y + Inches(1.2), card_w, Inches(0.7),
                    title, font_size=11, color=C_TEXT, align=PP_ALIGN.CENTER)

    # 특징 (하단)
    features = [
        ("📱", "PC · 모바일 반응형", "어디서든 접근"),
        ("🎯", "외국인 관점 재구성", "단순 번역 아님"),
        ("🔄", "오프라인 학습 가능", "캐시 활용"),
    ]
    for i, (icon, title, desc) in enumerate(features):
        x = Inches(0.7) + Inches(4.2) * i
        f_box = add_rounded(slide, x, Inches(4.7), Inches(3.9), Inches(1.6),
                             fill_color=C_WHITE, line_color=C_PRIMARY)
        add_textbox(slide, x + Inches(0.3), Inches(4.85), Inches(0.8), Inches(0.8),
                    icon, font_size=32, align=PP_ALIGN.LEFT)
        add_textbox(slide, x + Inches(1.2), Inches(4.95), Inches(2.6), Inches(0.5),
                    title, font_size=14, bold=True, color=C_DEEP)
        add_textbox(slide, x + Inches(1.2), Inches(5.4), Inches(2.6), Inches(0.5),
                    desc, font_size=12, color=C_GRAY)


def build_content_3_quiz(prs):
    """슬라이드 9 — 03 추진내용 ③ 4개 언어 1000문항 학습 사이트"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "03 추진내용 ③", "4개 언어 × 1000문항 학습 사이트 + 동영상", 9, TOTAL_PAGES)

    # 큰 수치
    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.6),
                "🌐  학과시험 문제은행을 4개 언어로 완벽 통일",
                font_size=20, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    # 메인 수치 박스
    main_box = add_rounded(slide, Inches(2), Inches(2.0), Inches(9), Inches(1.5),
                            fill_color=C_DEEP)
    add_textbox(slide, Inches(2), Inches(2.2), Inches(9), Inches(0.6),
                "4 언어 × 1000 문항 = 4,000 개 학습 자료",
                font_size=24, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(2), Inches(2.85), Inches(9), Inches(0.5),
                "한국어 · 영어 · 중국어 · 베트남어 · 모두 동일 분포",
                font_size=14, color=C_WHITE, align=PP_ALIGN.CENTER)

    # 카테고리 분포 4개 박스
    cats = [
        ("문장형", "680", C_PRIMARY),
        ("사진형", "185", C_ACCENT),
        ("표지판형", "100", C_SUCCESS),
        ("동영상형", "35", RGBColor(0xE6, 0x7E, 0x22)),
    ]
    for i, (label, num, color) in enumerate(cats):
        x = Inches(0.7) + Inches(3.1) * i
        c_box = add_rounded(slide, x, Inches(4.0), Inches(2.9), Inches(1.5),
                             fill_color=C_LIGHT_BG)
        add_textbox(slide, x, Inches(4.1), Inches(2.9), Inches(0.6),
                    label, font_size=15, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
        add_textbox(slide, x, Inches(4.65), Inches(2.9), Inches(0.85),
                    num, font_size=42, bold=True, color=color, align=PP_ALIGN.CENTER)

    # 동영상 강조
    add_textbox(slide, Inches(0.5), Inches(5.9), Inches(12.5), Inches(0.5),
                "🎬  동영상 35개 (Q966~Q1000) · GitHub Releases 호스팅 · 4개 언어 공통 사용",
                font_size=14, bold=True, color=C_PRIMARY, align=PP_ALIGN.CENTER)
    # 사이트 URL
    add_textbox(slide, Inches(0.5), Inches(6.4), Inches(12.5), Inches(0.5),
                "🔗  https://evergreenedu-collab.github.io/quiz/",
                font_size=12, color=C_GRAY, align=PP_ALIGN.CENTER)


def build_content_4_auto(prs):
    """슬라이드 10 — 03 추진내용 ④ 자동 갱신 시스템"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "03 추진내용 ④", "데이터 자동 갱신 시스템 — 지속가능성", 10, TOTAL_PAGES)

    # 핵심 메시지
    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.6),
                "💡  '한 번 만들고 끝'이 아닌, 매년 PDF 갱신 자동화",
                font_size=22, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.5), Inches(1.65), Inches(12.5), Inches(0.5),
                "사용자는 PDF만 다운로드 → AI 자동 처리 → 4개 언어 일괄 갱신",
                font_size=14, color=C_GRAY, align=PP_ALIGN.CENTER)

    # 갱신 워크플로 4단계 (가로 화살표)
    steps = [
        ("1", "PDF 다운로드", "사용자 작업"),
        ("2", "PyMuPDF 자동 추출", "1000문항 파싱"),
        ("3", "위치 매칭 알고리즘", "이미지 자동 연결"),
        ("4", "4언어 HTML 갱신", "PR 자동 생성"),
    ]
    step_w = Inches(2.6)
    step_h = Inches(2.0)
    step_y = Inches(2.6)
    for i, (no, title, desc) in enumerate(steps):
        x = Inches(0.4) + (step_w + Inches(0.4)) * i
        # 박스
        s_box = add_rounded(slide, x, step_y, step_w, step_h,
                             fill_color=C_LIGHT_BG, line_color=C_PRIMARY)
        # 번호 원
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.95),
                                          step_y + Inches(0.1), Inches(0.7), Inches(0.7))
        circle.shadow.inherit = False
        circle.fill.solid()
        circle.fill.fore_color.rgb = C_PRIMARY
        circle.line.fill.background()
        set_shape_text(circle, no, font_size=22, bold=True, color=C_WHITE)
        # 제목
        add_textbox(slide, x, step_y + Inches(0.95), step_w, Inches(0.5),
                    title, font_size=14, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
        # 설명
        add_textbox(slide, x, step_y + Inches(1.45), step_w, Inches(0.5),
                    desc, font_size=11, color=C_GRAY, align=PP_ALIGN.CENTER)
        # 화살표 (다음 박스로)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                             x + step_w + Inches(0.05),
                                             step_y + Inches(0.85),
                                             Inches(0.3), Inches(0.3))
            arrow.shadow.inherit = False
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = C_PRIMARY
            arrow.line.fill.background()

    # 하단 강조
    add_textbox(slide, Inches(0.5), Inches(5.2), Inches(12.5), Inches(0.5),
                "🛡️  CLAUDE.md · UPDATE_GUIDE.md · scripts/ 모두 GitHub에 영구 보관",
                font_size=14, bold=True, color=C_SUCCESS, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.5), Inches(5.8), Inches(12.5), Inches(0.6),
                "→ 향후 신규 PDF 배포 시 갱신 시간 ~ 10분 (기존 작업 시간 대폭 단축)",
                font_size=18, bold=True, color=C_PRIMARY, align=PP_ALIGN.CENTER)


def build_result_1_satisfaction(prs):
    """슬라이드 11 — 03 추진결과 ① 현장 만족도"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "03 추진결과 ①", "외국인 학습자 현장 만족도", 11, TOTAL_PAGES)

    # 거대 수치
    big_box = add_rounded(slide, Inches(2.5), Inches(1.5), Inches(8), Inches(2.7),
                           fill_color=C_DEEP)
    add_textbox(slide, Inches(2.5), Inches(1.65), Inches(8), Inches(0.5),
                "현장 만족도",
                font_size=18, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(2.5), Inches(2.2), Inches(8), Inches(1.5),
                "100%",
                font_size=120, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, font_name=FONT_EN)

    # 부속 메시지
    add_textbox(slide, Inches(0.5), Inches(4.5), Inches(12.5), Inches(0.7),
                "여러 기관 외국인 학습자 ○○○명 대상 강의 — 만족도 100%",
                font_size=22, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.5), Inches(5.3), Inches(12.5), Inches(0.5),
                "* 2026년 외국인 대상 교육 다수 예정 — 통합 만족도 조사 진행 중",
                font_size=12, color=C_GRAY, align=PP_ALIGN.CENTER)

    # 주요 평가 포인트 3개
    points = [
        ("자기 언어 학습", "모국어 질의응답"),
        ("실용성", "실제 시험 직결"),
        ("접근성", "PC · 모바일"),
    ]
    for i, (title, desc) in enumerate(points):
        x = Inches(0.7) + Inches(4.2) * i
        p_box = add_rounded(slide, x, Inches(6.0), Inches(3.9), Inches(1.0),
                             fill_color=C_LIGHT_BG, line_color=C_PRIMARY)
        add_textbox(slide, x, Inches(6.05), Inches(3.9), Inches(0.5),
                    title, font_size=14, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
        add_textbox(slide, x, Inches(6.55), Inches(3.9), Inches(0.4),
                    desc, font_size=11, color=C_GRAY, align=PP_ALIGN.CENTER)


def build_result_2_metrics(prs):
    """슬라이드 12 — 03 추진결과 ② 정량 성과"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "03 추진결과 ②", "정량 성과 지표", 12, TOTAL_PAGES)

    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.6),
                "🎯  AI 튜터 운영 핵심 지표",
                font_size=22, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    # 4개 큰 수치
    metrics = [
        ("4,000 +", "학습 자료", "4 언어 × 1000 문항", C_PRIMARY),
        ("35", "동영상", "966~1000번 문제", C_DEEP),
        ("24/7", "GPTs 운영", "365일 모국어 챗봇", C_ACCENT),
        ("100 %", "만족도", "외국인 학습자", C_SUCCESS),
    ]
    box_w = Inches(2.95)
    box_h = Inches(2.5)
    box_y = Inches(2.0)
    for i, (big, small, sub, color) in enumerate(metrics):
        x = Inches(0.5) + (box_w + Inches(0.2)) * i
        # 메인 박스
        add_rounded(slide, x, box_y, box_w, box_h, fill_color=C_LIGHT_BG)
        # 컬러 헤더
        add_rect(slide, x, box_y, box_w, Inches(0.4), fill_color=color)
        # 큰 수치
        add_textbox(slide, x, box_y + Inches(0.7), box_w, Inches(1.0),
                    big, font_size=44, bold=True, color=color, align=PP_ALIGN.CENTER)
        # 라벨
        add_textbox(slide, x, box_y + Inches(1.7), box_w, Inches(0.4),
                    small, font_size=14, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)
        # 설명
        add_textbox(slide, x, box_y + Inches(2.1), box_w, Inches(0.4),
                    sub, font_size=11, color=C_GRAY, align=PP_ALIGN.CENTER)

    # 부수 성과
    add_textbox(slide, Inches(0.5), Inches(4.95), Inches(12.5), Inches(0.5),
                "🛠️  부수 성과",
                font_size=16, bold=True, color=C_DEEP, align=PP_ALIGN.LEFT)

    extras = [
        "✓  4개 언어 데이터 정합성 100% 통일 (외국어별 누락·잉여 정정)",
        "✓  동영상 호스팅 자동화 (GitHub Releases 활용 · 무료 · 무제한 트래픽)",
        "✓  자동 갱신 도구 영구 보관 (다른 환경에서도 동일 작동)",
        "✓  외국인 사용 매뉴얼 PDF · QR 코드 제작",
    ]
    for i, item in enumerate(extras):
        add_textbox(slide, Inches(0.7), Inches(5.5 + i * 0.4), Inches(12), Inches(0.4),
                    item, font_size=13, color=C_TEXT)


def build_future_1(prs):
    """슬라이드 13 — 04 향후계획 ① 확장 가능성"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "04 향후계획 ①", "확장 가능성", 13, TOTAL_PAGES)

    add_textbox(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.7),
                "🚀  외국인 면허 종합 지원 플랫폼으로 확장",
                font_size=22, bold=True, color=C_DEEP, align=PP_ALIGN.CENTER)

    plans = [
        ("🌍", "추가 언어 확장",
         "몽골어 · 러시아어 · 필리핀어 · 태국어 등\n주요 외국인 거주 인구 언어 순차 추가",
         C_PRIMARY),
        ("📚", "학과시험 외 영역",
         "기능시험 · 도로주행 정보 통합 가이드,\n면허 갱신 · 적성검사 등 외국인 안내",
         C_ACCENT),
        ("🎙️", "AI 음성 기능",
         "음성 인식 · 합성으로 청각 · 시각 장애인,\n저학력 외국인 학습자도 접근 가능",
         C_SUCCESS),
    ]
    plan_y = Inches(2.0)
    plan_h = Inches(1.4)
    plan_gap = Inches(0.2)

    for i, (icon, title, desc, color) in enumerate(plans):
        y = plan_y + (plan_h + plan_gap) * i
        # 아이콘 박스
        ico_box = add_rounded(slide, Inches(0.7), y, Inches(1.5), plan_h,
                                fill_color=color)
        add_textbox(slide, Inches(0.7), y + Inches(0.3), Inches(1.5), Inches(0.8),
                    icon, font_size=44, align=PP_ALIGN.CENTER)
        # 본문
        add_rounded(slide, Inches(2.4), y, Inches(10.4), plan_h,
                     fill_color=C_LIGHT_BG)
        add_textbox(slide, Inches(2.7), y + Inches(0.2), Inches(10), Inches(0.5),
                    title, font_size=18, bold=True, color=C_DEEP)
        add_textbox(slide, Inches(2.7), y + Inches(0.7), Inches(10), Inches(0.7),
                    desc, font_size=12, color=C_TEXT)


def build_future_2(prs):
    """슬라이드 14 — 04 향후계획 ② 본부 모범사례 잠재력"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "04 향후계획 ②", "본부 모범사례 · 정책 부합성", 14, TOTAL_PAGES)

    # 큰 비전
    vision_box = add_rounded(slide, Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.2),
                              fill_color=C_DEEP)
    set_shape_text(vision_box,
                   "부산지부의 시작 → 전국 공단 표준 모델로",
                   font_size=24, bold=True, color=C_WHITE)

    items = [
        ("🏆", "본부 모범사례 잠재력",
         "외국인 정착 지원이라는 사회적 가치 + AI 활용 혁신성"),
        ("🇰🇷", "정책적 의의",
         "디지털 정부 · AI 정부 · 다문화 사회 정책 부합"),
        ("📈", "교통안전 향상",
         "외국인 운전자 학과시험 합격률 ↑ → 도로 안전 ↑"),
        ("🌐", "전국 확산",
         "타 지부에 동일 도구 · 절차 그대로 이식 가능 (자동 갱신 시스템)"),
    ]
    for i, (icon, title, desc) in enumerate(items):
        x = Inches(0.7) if i % 2 == 0 else Inches(7.0)
        y = Inches(2.9) if i < 2 else Inches(4.7)
        item_box = add_rounded(slide, x, y, Inches(5.9), Inches(1.6),
                                fill_color=C_LIGHT_BG, line_color=C_PRIMARY)
        add_textbox(slide, x + Inches(0.2), y + Inches(0.2), Inches(0.8), Inches(1.2),
                    icon, font_size=36, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(1.1), y + Inches(0.2), Inches(4.6), Inches(0.5),
                    title, font_size=16, bold=True, color=C_DEEP)
        add_textbox(slide, x + Inches(1.1), y + Inches(0.75), Inches(4.6), Inches(0.8),
                    desc, font_size=12, color=C_TEXT)

    # 마무리 메시지
    add_textbox(slide, Inches(0.5), Inches(6.7), Inches(12.5), Inches(0.5),
                "→ 외국인 정착 지원의 새 표준을 부산지부가 선도한다",
                font_size=16, bold=True, color=C_PRIMARY, align=PP_ALIGN.CENTER)


def build_closing(prs):
    """슬라이드 15 — 맺음말"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # 배경 띠
    add_rect(slide, Inches(0), Inches(0), Inches(0.6), Inches(7.5), fill_color=C_DEEP)
    add_rect(slide, Inches(0.6), Inches(0), Inches(0.15), Inches(7.5), fill_color=C_PRIMARY)

    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(10.5), Inches(0.4), width=Inches(2.5))

    # 메인 메시지
    add_textbox(slide, Inches(1.5), Inches(2.0), Inches(11), Inches(1.0),
                "외국인이 자신의 모국어로",
                font_size=36, color=C_TEXT, align=PP_ALIGN.LEFT)
    add_textbox(slide, Inches(1.5), Inches(2.7), Inches(11), Inches(1.0),
                "한국 운전면허를 배우는 시대",
                font_size=44, bold=True, color=C_DEEP, align=PP_ALIGN.LEFT)

    # 강조 박스
    box = add_rounded(slide, Inches(1.5), Inches(4.5), Inches(11), Inches(1.0),
                       fill_color=C_PRIMARY)
    set_shape_text(box,
                   "AI · 다국어 · 자동화로 외국인 정착의 첫걸음을 함께",
                   font_size=20, bold=True, color=C_WHITE)

    # 감사
    add_textbox(slide, Inches(1.5), Inches(6.0), Inches(11), Inches(0.8),
                "감사합니다.",
                font_size=32, bold=True, color=C_DEEP)

    # 부서명
    add_textbox(slide, Inches(1.5), Inches(6.7), Inches(11), Inches(0.5),
                "한국도로교통공단 부산광역시지부 안전교육부",
                font_size=14, color=C_GRAY)


# ----- 메인 -----

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=None, help="출력 PPT 경로")
    args = parser.parse_args()

    if args.out:
        out_path = Path(args.out)
    else:
        desktop = Path(r"C:\Users\user\OneDrive\바탕 화면")
        date_tag = datetime.now().strftime("%Y-%m")
        out_path = desktop / f"외국인_학과시험_AI튜터_보고서_{date_tag}.pptx"

    out_path.parent.mkdir(parents=True, exist_ok=True)

    # 16:9 와이드 (13.333 × 7.5 인치)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    print("=" * 60)
    print("외국인 학과시험 AI 튜터 — 보고서 PPT 빌드")
    print("=" * 60)

    builders = [
        ("표지", build_cover),
        ("기본현황 / 4대 통합", build_overview),
        ("01 추진배경 ① 외국인 증가", build_bg_1),
        ("01 추진배경 ② 3가지 어려움", build_bg_2),
        ("01 추진배경 ③ 공공 형평성", build_bg_3),
        ("02 추진계획 4대 축", build_plan),
        ("03 추진내용 ① GPTs", build_content_1_gpts),
        ("03 추진내용 ② 교재", build_content_2_textbook),
        ("03 추진내용 ③ 학습 사이트", build_content_3_quiz),
        ("03 추진내용 ④ 자동 갱신", build_content_4_auto),
        ("03 추진결과 ① 만족도", build_result_1_satisfaction),
        ("03 추진결과 ② 정량성과", build_result_2_metrics),
        ("04 향후계획 ① 확장", build_future_1),
        ("04 향후계획 ② 모범사례", build_future_2),
        ("맺음말", build_closing),
    ]

    for i, (name, fn) in enumerate(builders, 1):
        fn(prs)
        print(f"  {i:2d}/15  {name}")

    prs.save(str(out_path))
    print()
    print(f"✅ 저장 완료: {out_path}")
    print(f"   파일 크기: {out_path.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
